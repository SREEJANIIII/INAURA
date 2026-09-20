import io

from app.services import document_parser as dp
from app.services import signal_extractor as se


def _make_docx_bytes(lines: list[str]) -> bytes:
    docx = pytest_import_docx()
    doc = docx.Document()
    for line in lines:
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def pytest_import_docx():
    import docx

    return docx


def test_split_sections_detects_headers():
    text = "John Doe\nSKILLS\nPython, React\nEXPERIENCE\nBuilt APIs with Python"
    sections = dp.split_sections(text)
    assert sections["skills"] == "Python, React"
    assert "Built APIs" in sections["experience"]


def test_parse_docx_extracts_text_and_sections():
    data = _make_docx_bytes(
        [
            "John Doe Resume",
            "SKILLS",
            "Python, React, PostgreSQL",
            "EXPERIENCE",
            "Engineered scalable microservices using Python, FastAPI, and PostgreSQL.",
        ]
    )
    meta = dp.parse_document_bytes(
        data,
        filename="resume.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert meta["parse_status"] == "ok"
    assert "Python" in meta["parsed_text"]
    assert meta["word_count"] > 5
    assert "skills" in meta["sections"]
    assert "experience" in meta["sections"]


def test_parse_pdf_fallback_regex():
    pdf = (
        b"%PDF-1.4\n1 0 obj << >> endobj\n"
        b"BT /F1 12 Tf 72 720 Td (Engineered microservices using Python and PostgreSQL) Tj ET\n"
        b"%%EOF"
    )
    meta = dp.parse_document_bytes(pdf, filename="resume.pdf", content_type="application/pdf")
    assert meta["parse_status"] == "ok"
    assert "Python" in meta["parsed_text"]


def test_parse_empty_never_raises():
    meta = dp.parse_document_bytes(b"", filename="resume.pdf", content_type="application/pdf")
    assert meta["parsed_text"] == ""
    assert meta["parse_status"] == "empty"


def test_parse_garbage_never_raises():
    meta = dp.parse_document_bytes(b"\x00\x01\x02not-a-doc", filename="resume.pdf")
    assert meta["parsed_text"] == ""
    assert meta["parse_status"] in ("failed", "empty", "ok")


def test_parsed_resume_flows_into_signals():
    """End-to-end: parsed docx metadata must produce skill signals (the pre-fix gap)."""
    data = _make_docx_bytes(
        [
            "Jane Doe",
            "EXPERIENCE",
            "Engineered scalable microservices using Python, FastAPI, and PostgreSQL. "
            "Implemented responsive frontend interfaces with React.",
        ]
    )
    meta = dp.parse_document_bytes(data, filename="resume.docx")
    assert meta["parse_status"] == "ok"

    evidence = [{"id": "doc1", "evidence_type": "resume", "title": "resume.docx", "metadata": meta}]
    signals = se.extract_signals(evidence, [], [])
    detected = {s["canonical_name"] for s in signals}
    assert "Python" in detected
    assert "React" in detected
    assert "PostgreSQL" in detected
    # Concrete action verbs -> 0.50 action strength
    py = next(s for s in signals if s["canonical_name"] == "Python")
    assert py["signal_value"] == 0.50
    assert py["source_reliability"] == 0.50


def test_clean_text_preserves_urls_and_emails():
    raw = (
        "Contact: alex.Morgan@example.com or mailto:alex@gmail.com\n"
        "Profile: https://github.com/alex?tab=repos&filter=all\n"
        "Skills:Python,FastAPI,PostgreSQL.Built microservices.He led teams."
    )
    cleaned = dp._clean_text(raw)
    assert "alex.Morgan@example.com" in cleaned
    assert "mailto:alex@gmail.com" in cleaned
    assert "https://github.com/alex?tab=repos&filter=all" in cleaned
    assert "Python, FastAPI, PostgreSQL." in cleaned
    assert "Built microservices. He led teams." in cleaned


def test_split_sections_numbered_and_multicolumn():
    text = (
        "Alex Morgan\n"
        "alex@example.com\n"
        "1. PROFESSIONAL EXPERIENCE\n"
        "Software Engineer at Acme Corp\n"
        "2. TECHNICAL SKILLS\n"
        "Python, React, AWS\n"
        "Section 3: EDUCATION\n"
        "B.S. in Computer Science"
    )
    sections = dp.split_sections(text)
    assert "experience" in sections
    assert "skills" in sections
    assert "education" in sections
    assert "Software Engineer" in sections["experience"]
    assert "Python" in sections["skills"]
    assert "B.S. in Computer Science" in sections["education"]


def test_split_sections_dual_header_line():
    text = (
        "Alex Morgan\n"
        "alex@example.com\n"
        "PROFESSIONAL EXPERIENCE           TECHNICAL SKILLS\n"
        "Engineered scalable backends with Python."
    )
    sections = dp.split_sections(text)
    assert "experience" in sections
    assert "skills" in sections


def test_docx_header_and_rels_extraction():
    docx = pytest_import_docx()
    doc = docx.Document()
    # Put contact info inside DOCX header
    header = doc.sections[0].header
    hp = header.paragraphs[0]
    hp.text = "Alex Morgan | alex.morgan@gmail.com | github.com/alexmorgan"

    # Put experience in body
    doc.add_paragraph("WORK EXPERIENCE")
    doc.add_paragraph("Engineered Python backends.")

    buf = io.BytesIO()
    doc.save(buf)

    meta = dp.parse_document_bytes(buf.getvalue(), filename="resume.docx")
    assert meta["parse_status"] == "ok"
    assert "alex.morgan@gmail.com" in meta["parsed_text"]
    assert "Alex Morgan" in meta["parsed_text"]
    assert any("github.com/alexmorgan" in link for link in meta["links"])
    assert "experience" in meta["sections"]


def test_pdf_annotation_links_extraction():
    import pypdf
    from pypdf.generic import DictionaryObject, NameObject, TextStringObject, ArrayObject, NumberObject

    writer = pypdf.PdfWriter()
    page = writer.add_blank_page(width=612, height=792)

    # Add link annotation
    link_annot = DictionaryObject()
    link_annot[NameObject("/Type")] = NameObject("/Annot")
    link_annot[NameObject("/Subtype")] = NameObject("/Link")
    link_annot[NameObject("/Rect")] = ArrayObject([NumberObject(100), NumberObject(700), NumberObject(200), NumberObject(720)])
    action = DictionaryObject()
    action[NameObject("/S")] = NameObject("/URI")
    action[NameObject("/URI")] = TextStringObject("www.linkedin.com/in/alexmorgan")
    link_annot[NameObject("/A")] = action
    writer.add_annotation(page_number=0, annotation=link_annot)

    buf = io.BytesIO()
    writer.write(buf)

    # Also test extracting annotations directly
    reader = pypdf.PdfReader(buf)
    extracted = dp._extract_pdf_annotations(reader.pages[0])
    assert "https://www.linkedin.com/in/alexmorgan" in extracted


def test_docx_textbox_contact_info_extracted():
    docx = pytest_import_docx()
    from docx.oxml import parse_xml

    doc = docx.Document()
    # Add a paragraph with a textbox containing email and phone
    p = doc.add_paragraph()
    xml_textbox = parse_xml(
        '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '  <w:drawing>'
        '    <w:txbxContent>'
        '      <w:p><w:r><w:t>Email: candidate@techcorp.dev | Phone: +91 98765 43210</w:t></w:r></w:p>'
        '    </w:txbxContent>'
        '  </w:drawing>'
        '</w:r>'
    )
    p._p.append(xml_textbox)
    doc.add_paragraph("WORK EXPERIENCE")
    doc.add_paragraph("Engineered scalable backends with Python.")

    buf = io.BytesIO()
    doc.save(buf)

    meta = dp.parse_document_bytes(buf.getvalue(), filename="resume.docx")
    assert meta["parse_status"] == "ok"
    assert "candidate@techcorp.dev" in meta["parsed_text"]
    assert "98765 43210" in meta["parsed_text"]
    assert "experience" in meta["sections"]


